"""Read supplied workbooks without executing formulas, VBA or external links."""
import json, hashlib, re, sys
from pathlib import Path
from collections import Counter
import openpyxl
from openpyxl.utils import get_column_letter as col
SRC=Path('C:/Users/John/Desktop/DD automations/Aircraft Noise Automation')
OUT=Path('outputs/aircraft-noise/aircraft_noise/data')
E='aircraft_noise_extraction2.xlsx'; W='AS2021-2015 Aircraft Noise 1.01.xlsm'
ext=openpyxl.load_workbook(SRC/E,read_only=True,data_only=False,keep_links=False)
old=openpyxl.load_workbook(SRC/W,read_only=True,data_only=False,keep_links=False)
# Materialise once; never invoke Excel/VBA or trust cached formula results.
mat=list(ext['Aircraft Matrices'].iter_rows(values_only=True))
index={r[0]:dict(table_id=r[0],title=r[1],aircraft=r[2],operation=r[3],qualifier=r[4],pdf_page=r[5],printed_page=r[6]) for r in ext['Table Index'].iter_rows(min_row=4,values_only=True) if r[0]}
tables={}; problems=[]; comparisons=Counter(); diffs=[]
old_values={name:list(old[name].iter_rows(values_only=True)) for name in old.sheetnames if re.fullmatch(r'3\.\d+\([ABC]\)',name)}
for tid,info in index.items():
    suffix=tid[-2]; idcol={'A':0,'B':1,'C':2}[suffix]; start,end={'A':(5,26),'B':(26,45),'C':(45,64)}[suffix]
    rows=[(i+1,r) for i,r in enumerate(mat[3:],start=3) if r[idcol]==tid]
    inferred=False
    if not rows and suffix in ['B','C']:
        # Five B identifiers are blank. Wide columns and Table Index disambiguate
        # the populated departure block belonging to the A-labelled aircraft.
        rows=[(i+1,r) for i,r in enumerate(mat[3:],start=3) if r[0]==tid[:-3]+'(A)' and any(v is not None for v in r[start:end])]
        inferred=bool(rows)
    ds=list(mat[2][start:end]); dist=[r[4] for _,r in rows]; values=[list(r[start:end]) for _,r in rows]
    cells=[[f'{col(c+1)}{i}' for c in range(start,end)] for i,r in rows]
    if not rows or dist != sorted(set(dist)): problems.append({'table':tid,'problem':'Missing rows or invalid distance axis'})
    table={**info,'distance_axis_m':dist,'ds_axis_m':ds,'values_db':values,'cell_refs':cells,
           'source_sheet':'Aircraft Matrices','source_refs':['noise-extraction'],'verification_status':'requires_confirmation',
           'identifier_recovery':'Blank departure ID: matched A block, global departure column headers and Table Index' if inferred else None}
    previous=old_values.get(tid)
    if previous:
        old_ds=previous[1][1:1+len(ds)];old_rows={r[0]:r for r in previous[2:] if isinstance(r[0],(int,float))}
        for i,d in enumerate(dist):
            for j,s in enumerate(ds):
                if s not in old_ds or d not in old_rows:comparisons['missing_axis']+=1;continue
                oi=old_ds.index(s)+1; v=old_rows[d][oi]; nv=values[i][j]
                comparisons['compared']+=1
                if v==nv:comparisons['match']+=1
                else:
                    comparisons['different']+=1
                    diffs.append({'table':tid,'distance_m':d,'ds_m':s,'extraction_cell':cells[i][j],'extraction_value':nv,'workbook_value':v})
    tables[tid]=table
freq=[]; ws=old['Aircraft Type Frequency ANEF -2']; arr=list(ws.iter_rows(values_only=True))
for row in range(7,26):
    vals=arr[row-1]; code=vals[1]
    for runway,start in [('07',2),('25',9),('16L',16),('34R',23),('16R',30),('34L',37)]:
        for op,offset in [('arrival',0),('departure',2)]:
            a,b=vals[start+offset:start+offset+2]
            numeric=[v for v in [a,b] if isinstance(v,(float,int))]
            freq.append({'aircraft_code':code,'runway_direction':runway,'operation':op,'day_raw':a,'night_raw':b,
                         'known_movement_sum':sum(numeric) if numeric else None,'movement_unit':'not stated on supplied worksheet',
                         'has_positive_movement':any(v>0 for v in numeric),'has_unresolved_marker':len(numeric)<2,
                         'source_sheet':ws.title,'source_cells':[f'{col(start+offset+1)}{row}',f'{col(start+offset+2)}{row}'],'source_refs':['frequency-workbook']})
mapping=[]
for i,r in enumerate(ext['Aircraft Mapping'].iter_rows(min_row=4,values_only=True),start=4):
    if r[0]:mapping.append({'group':r[0],'manufacturer':r[1],'model':r[2],'table':r[3],'representative_aircraft':r[4],'source_cells':f'C{i}:E{i}'})
heights=[list(r[:5]) for r in ext['Land Height'].iter_rows(min_row=4,values_only=True) if isinstance(r[0],(int,float))]
sources=[]
for name,sid,kind in [(E,'noise-extraction','extracted_reference_tables'),(W,'frequency-workbook','forecast_workbook')]:
    sources.append({'source_id':sid,'title':name,'filename':name,'sha256':hashlib.sha256((SRC/name).read_bytes()).hexdigest(),
                    'edition_or_revision':'AS 2021:2015 / supplied 2026-09-26','source_type':kind,
                    'locator':{'method':'openpyxl read-only raw values/formulas; no macros run','sheets':['Aircraft Matrices','Table Index','Aircraft Mapping','Land Height'] if sid=='noise-extraction' else ['Aircraft Type Frequency ANEF -2','Draft Sydney Aircraft Noise'],
                               'verification':'cross-workbook reconciliation, not independent original-page verification'},'verification_status':'requires_confirmation'})
ref={'edition':'AS 2021:2015','forecast_year':2045,'tables':tables,'frequency':freq,'mapping':mapping,'height_rows':heights,'sources':sources,
     'audit':{'matrix_count':len(tables),'frequency_sheet':ws.title,'frequency_codes':list(dict.fromkeys(x['aircraft_code'] for x in freq)),
              'comparison':dict(comparisons),'differences':diffs,'structural_issues':problems,'workbook_lookup_formula':old_values['3.4(A)'][3][24],
              'ranking_formula':"'Draft Sydney Aircraft Noise'!E10:E12: LARGE(D42:D71,1/2/3)",
              'frequency_units':'Worksheet does not state units or day/night time definitions; fractional raw values retained; no rate threshold used.'}}
(OUT/'reference_data.json').write_text(json.dumps(ref,ensure_ascii=False,allow_nan=False),encoding='utf-8')
print(json.dumps({'tables':len(tables),'frequency_rows':len(freq),'comparison':dict(comparisons),'structural_issues':problems,'sample_differences':diffs[:12]},ensure_ascii=True,indent=2))
