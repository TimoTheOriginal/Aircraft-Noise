import json
from pathlib import Path
import openpyxl
src=Path('C:/Users/John/Desktop/DD automations/Aircraft Noise Automation')
for name in ['aircraft_noise_extraction2.xlsx','AS2021-2015 Aircraft Noise 1.01.xlsm']:
    wb=openpyxl.load_workbook(src/name,read_only=True,data_only=False,keep_links=False)
    out={}
    print('\nWORKBOOK',name)
    for ws in wb:
        rows=[]
        for row in ws:
            vals={c.coordinate:c.value for c in row if c.value is not None}
            if vals: rows.append(vals)
        out[ws.title]=rows
        print(ws.title, ws.max_row, ws.max_column, 'nonempty rows',len(rows))
        print(json.dumps(rows[:5],ensure_ascii=False,default=str)[:1800])
    Path('work/'+name+'.json').write_text(json.dumps(out,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
